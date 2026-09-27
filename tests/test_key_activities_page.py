from datetime import date, timedelta

from PySide6.QtCore import QDate
from PySide6.QtWidgets import QFileDialog, QMessageBox

from database.models import (KeyActivity, KEY_STATUS_DONE, KEY_STATUS_IN_PROGRESS,
                             KEY_PRIORITY_TOP)
from ui.key_activities_page import KeyActivitiesPage


def _add(db, tech, title, **kw):
    kw.setdefault("start_date", date.today())
    kw.setdefault("status", KEY_STATUS_IN_PROGRESS)
    kw.setdefault("progress", 0)
    act = KeyActivity(technician_id=tech.id, technician_name_snapshot=tech.full_name,
                      title=title, **kw)
    db.add(act)
    db.commit()
    return act


def _page(db, tech, admin=False):
    page = KeyActivitiesPage(db, tech, admin=admin)
    page.period.set_key("all")
    return page


def test_default_period_is_this_month(qapp, db, owner, dialogs):
    assert KeyActivitiesPage(db, owner).period.key() == "this_month"


def test_owner_sees_only_own_activities(qapp, db, owner, other, dialogs):
    _add(db, owner, "A")
    _add(db, owner, "B")
    _add(db, other, "C")
    page = _page(db, owner)
    assert page.table.rowCount() == 2
    assert {a.title for a in page.shown} == {"A", "B"}
    assert page.btn_new is not None and page.update_box is not None
    assert page.cmb_owner is None


def test_admin_sees_all_read_only_with_owner_filter(qapp, db, owner, other, admin, dialogs):
    _add(db, owner, "A")
    _add(db, owner, "B")
    _add(db, other, "C")
    page = _page(db, admin, admin=True)
    assert page.table.rowCount() == 3
    assert page.btn_new is None and page.update_box is None
    page.cmb_owner.setCurrentIndex(page.cmb_owner.findData(other.id))
    assert [a.title for a in page.shown] == ["C"]


def test_selecting_a_row_shows_details(qapp, db, owner, dialogs):
    _add(db, owner, "سرور بکاپ", description="بکاپ شبانه")
    page = _page(db, owner)
    page.table.selectRow(0)
    assert page.current is not None and page.current.title == "سرور بکاپ"
    assert page.update_box.isEnabled()


def test_add_update_from_detail_panel(qapp, db, owner, dialogs):
    act = _add(db, owner, "سرور بکاپ", progress=10)
    page = _page(db, owner)
    page.table.selectRow(0)
    page.spn_update_progress.setValue(60)
    page.txt_update.setText("استوریج نصب شد")
    page.btn_add_update.click()
    db.refresh(act)
    assert act.progress == 60
    assert len(act.updates) == 1 and act.updates[0].text == "استوریج نصب شد"
    assert act.updates[0].author_name_snapshot == owner.full_name
    assert page.txt_update.text() == ""
    assert page.current is not None and page.current.id == act.id
    assert page.lst_timeline.count() == 1


def test_empty_update_text_is_rejected(qapp, db, owner, dialogs):
    calls, _ = dialogs
    act = _add(db, owner, "سرور بکاپ")
    page = _page(db, owner)
    page.table.selectRow(0)
    page.btn_add_update.click()
    db.refresh(act)
    assert act.updates == []
    assert calls[-1][0] == "warning"


def test_full_progress_offers_done_status(qapp, db, owner, dialogs):
    calls, answers = dialogs
    answers["question"] = QMessageBox.Yes
    act = _add(db, owner, "سرور بکاپ", progress=80)
    page = _page(db, owner)
    page.table.selectRow(0)
    page.spn_update_progress.setValue(100)
    page.txt_update.setText("تحویل شد")
    page.btn_add_update.click()
    db.refresh(act)
    assert any(kind == "question" for kind, _ in calls)
    assert act.status == KEY_STATUS_DONE
    assert act.end_date == date.today()


def test_full_progress_declined_keeps_status(qapp, db, owner, dialogs):
    _, answers = dialogs
    answers["question"] = QMessageBox.No
    act = _add(db, owner, "سرور بکاپ", progress=80)
    page = _page(db, owner)
    page.table.selectRow(0)
    page.spn_update_progress.setValue(100)
    page.txt_update.setText("تقریباً تمام")
    page.btn_add_update.click()
    db.refresh(act)
    assert act.progress == 100
    assert act.status == KEY_STATUS_IN_PROGRESS


def test_summary_strip_counts(qapp, db, owner, dialogs):
    _add(db, owner, "done", status=KEY_STATUS_DONE, progress=100, end_date=date.today())
    _add(db, owner, "open")
    page = _page(db, owner)
    assert page.stat_labels["done_in_period"].text() == "1"
    assert page.stat_labels["in_progress"].text() == "1"
    assert page.stat_labels["started_in_period"].text() == "2"
    assert page.stat_labels["updates_in_period"].text() == "0"


def test_filters_and_search(qapp, db, owner, dialogs):
    _add(db, owner, "رفع قطعی شبکه", category="شبکه", priority=KEY_PRIORITY_TOP)
    _add(db, owner, "خرید لایسنس", category="نرم‌افزار")
    page = _page(db, owner)
    page.cmb_category.setCurrentIndex(page.cmb_category.findData("شبکه"))
    assert [a.title for a in page.shown] == ["رفع قطعی شبکه"]
    page.cmb_category.setCurrentIndex(0)
    page.txt_search.setText("لایسنس")
    assert [a.title for a in page.shown] == ["خرید لایسنس"]


def test_delete_activity_with_confirmation(qapp, db, owner, dialogs):
    _add(db, owner, "حذفی")
    page = _page(db, owner)
    page.table.selectRow(0)
    page.btn_delete.click()
    assert db.query(KeyActivity).count() == 0
    assert page.table.rowCount() == 0
    assert page.current is None


def test_delete_cancelled_keeps_activity(qapp, db, owner, dialogs):
    _, answers = dialogs
    answers["question"] = QMessageBox.No
    _add(db, owner, "بماند")
    page = _page(db, owner)
    page.table.selectRow(0)
    page.btn_delete.click()
    assert db.query(KeyActivity).count() == 1


def test_export_writes_file(qapp, db, owner, dialogs, monkeypatch, tmp_path):
    _add(db, owner, "A")
    target = tmp_path / "out.xlsx"
    monkeypatch.setattr(QFileDialog, "getSaveFileName", lambda *a, **k: (str(target), ""))
    page = _page(db, owner)
    page.export()
    assert target.exists()


def test_export_with_nothing_shown_warns(qapp, db, owner, dialogs):
    calls, _ = dialogs
    page = _page(db, owner)
    page.export()
    assert calls[-1][0] == "warning"


def test_admin_export_with_owner_filter_names_that_owner(qapp, db, owner, other, admin,
                                                          dialogs, monkeypatch, tmp_path):
    import openpyxl
    _add(db, owner, "A")
    _add(db, other, "C")
    target = tmp_path / "out.xlsx"
    monkeypatch.setattr(QFileDialog, "getSaveFileName", lambda *a, **k: (str(target), ""))
    page = _page(db, admin, admin=True)
    page.cmb_owner.setCurrentIndex(page.cmb_owner.findData(other.id))
    page.export()
    wb = openpyxl.load_workbook(target)
    ws = wb["جمع‌بندی"]
    summary = {ws.cell(row=r, column=1).value: ws.cell(row=r, column=2).value
              for r in range(2, ws.max_row + 1)}
    assert summary["ثبت‌کننده"] == other.full_name


def _row_of(page, act):
    return next(i for i, a in enumerate(page.shown) if a.id == act.id)


def _delete_elsewhere(activity_id):
    """شبیه‌سازیِ حذف همان ردیف از یک اتصال دیگر (مثلاً یک ماشین دیگر)."""
    from database.connection import SessionLocal
    other_db = SessionLocal()
    row = other_db.get(KeyActivity, activity_id)
    other_db.delete(row)
    other_db.commit()
    other_db.close()


def test_load_data_survives_deleted_selected_activity(qapp, db, owner, dialogs):
    act = _add(db, owner, "کار موقت")
    page = _page(db, owner)
    page.table.selectRow(0)
    assert page.current is not None
    _delete_elsewhere(act.id)
    db.expire_all()   # شبیه‌سازیِ اکسپایر شدن شیء پس از commit روی سشن مشترک
    page.load_data()  # نباید ObjectDeletedError بدهد
    assert page.current is None
    assert page.table.rowCount() == 0


def test_apply_filters_after_rollback_survives_deleted_other_activity(qapp, db, owner, dialogs):
    """رول‌بکِ خطای یک کار نباید کارِ دیگری را که هم‌زمان جای دیگر حذف شده، در
    self.all اکسپایرشده نگه دارد؛ وگرنه فیلتر/جستجوی بعدی با ObjectDeletedError می‌ترکد."""
    act_a = _add(db, owner, "الف", start_date=date.today())
    act_b = _add(db, owner, "ب")
    page = _page(db, owner)
    page.table.selectRow(_row_of(page, act_a))
    assert page.current is not None and page.current.id == act_a.id

    _delete_elsewhere(act_b.id)
    db.expire_all()

    # مسیر ValueError در add_update: تاریخِ به‌روزرسانی قبل از تاریخ شروعِ کار ← رول‌بک
    before_start = act_a.start_date - timedelta(days=1)
    page.date_update.setDate(QDate(before_start.year, before_start.month, before_start.day))
    page.txt_update.setText("متن به‌روزرسانی")
    page.btn_add_update.click()

    calls, _ = dialogs
    assert any(kind == "warning" for kind, _ in calls)

    # نباید ObjectDeletedError بدهد (self.all دیگر نباید کارِ حذف‌شده‌ی «ب» را نگه دارد)
    page.txt_search.setText("x")
    page.txt_search.setText("")
    assert page.table.rowCount() == 1


def test_edit_activity_handles_deleted_selection(qapp, db, owner, dialogs):
    calls, _ = dialogs
    act = _add(db, owner, "کار موقت")
    page = _page(db, owner)
    page.table.selectRow(0)
    _delete_elsewhere(act.id)
    db.expire_all()
    page.edit_activity()  # نباید ObjectDeletedError بدهد
    assert any(kind == "information" for kind, _ in calls)
    assert page.current is None
    assert page.table.rowCount() == 0
