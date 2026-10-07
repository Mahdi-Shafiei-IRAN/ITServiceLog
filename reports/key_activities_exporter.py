"""خروجی اکسل «کارهای شاخص IT» — سه شیت، هم‌سبک با بقیه‌ی خروجی‌های برنامه.

  1. «کارهای شاخص»  — هر ردیف یک کار با وضعیت فعلی‌اش
  2. «گزارش پیشرفت» — فقط به‌روزرسانی‌های داخل بازه («در این بازه چه شد»)
  3. «جمع‌بندی»     — همان عددهای نوار خلاصه‌ی برنامه + تفکیک‌ها
"""
import openpyxl
from openpyxl.utils import get_column_letter

from reports.excel_exporter import _new_sheet, _write_row
from services import key_activity_service as svc
from utils import jalali


def _dt(d):
    return jalali.fmt(d)


def range_label(rng):
    return "همه‌ی زمان‌ها" if rng is None else f"{_dt(rng[0])} تا {_dt(rng[1])}"


def default_filename(rng):
    if rng is None:
        return "Key_Activities_All.xlsx"
    return f"Key_Activities_{jalali.fmt(rng[0], '-')}_{jalali.fmt(rng[1], '-')}.xlsx"


def export_key_activities_to_excel(filepath, activities, rng, owner=None, show_owner=False):
    wb = openpyxl.Workbook()

    # ---------------- شیت ۱: کارهای شاخص ----------------
    headers = ["ردیف", "عنوان", "دسته", "اهمیت", "وضعیت", "پیشرفت ٪", "شروع", "اتمام"]
    if show_owner:
        headers.append("ثبت‌کننده")
    headers += ["شرح", "نتیجه"]
    ws = _new_sheet(wb, "کارهای شاخص", headers, first=True)
    for i, a in enumerate(activities, 1):
        values = [i, a.title, a.category or "-", a.priority or "-", a.status or "-",
                  a.progress or 0, _dt(a.start_date), _dt(a.end_date)]
        if show_owner:
            values.append(a.technician_name_snapshot or "-")
        values += [a.description or "-", a.result or "-"]
        _write_row(ws, i + 1, values, wrap_from=2)
    ws.column_dimensions["B"].width = 40
    ws.column_dimensions["C"].width = 16
    ws.column_dimensions["E"].width = 14
    if show_owner:
        ws.column_dimensions["I"].width = 18
    for col in (len(headers) - 1, len(headers)):  # شرح و نتیجه
        ws.column_dimensions[get_column_letter(col)].width = 45

    # ---------------- شیت ۲: گزارش پیشرفت ----------------
    ws2 = _new_sheet(wb, "گزارش پیشرفت", [
        "ردیف", "تاریخ", "عنوان کار", "متن به‌روزرسانی", "پیشرفت ٪", "وضعیت", "ثبت‌کننده",
    ])
    for i, u in enumerate(svc.updates_in_period(activities, rng), 1):
        _write_row(ws2, i + 1, [
            i,
            _dt(u.update_date),
            u.activity.title if u.activity else "-",
            u.text or "-",
            u.progress if u.progress is not None else "-",
            u.status or "-",
            u.author_name_snapshot or "-",
        ], wrap_from=3)
    ws2.column_dimensions["C"].width = 36
    ws2.column_dimensions["D"].width = 55
    ws2.column_dimensions["G"].width = 18

    # ---------------- شیت ۳: جمع‌بندی ----------------
    s = svc.summarize(activities, rng)
    rows = []
    if owner is not None:
        rows.append(("ثبت‌کننده", owner.full_name))
    rows.append(("بازه‌ی گزارش", range_label(rng)))
    rows += [
        ("تعداد کارهای شاخص در بازه", s["total"]),
        ("انجام‌شده در بازه", s["done_in_period"]),
        ("شروع‌شده در بازه", s["started_in_period"]),
        ("در حال انجام", s["in_progress"]),
        ("تعداد به‌روزرسانی‌ها در بازه", s["updates_in_period"]),
    ]
    for title, key in (("به تفکیک وضعیت", "by_status"),
                       ("به تفکیک دسته", "by_category"),
                       ("به تفکیک اهمیت", "by_priority")):
        rows.append(("", ""))
        rows.append((title, ""))
        rows.extend(sorted(s[key].items(), key=lambda kv: kv[1], reverse=True))
    ws3 = _new_sheet(wb, "جمع‌بندی", ["عنوان", "مقدار"])
    for i, (key, value) in enumerate(rows, 2):
        _write_row(ws3, i, [key, value], wrap_from=1)
    ws3.column_dimensions["A"].width = 40
    ws3.column_dimensions["B"].width = 24

    wb.save(filepath)
